"""
core/loader.py
--------------
Reads an uploaded file and turns it into LangChain `Document` objects,
exactly like the Document(page_content=..., metadata=...) cells in Simple_RAG.ipynb.

Supported: .pdf  .txt  .md  .csv  .docx
"""

import os
from langchain_core.documents import Document


def _read_pdf(path, filename):
    from pypdf import PdfReader
    reader = PdfReader(path)
    docs = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or '').strip()
        if text:
            docs.append(Document(page_content=text,
                                 metadata={'source': filename, 'page': i}))
    return docs


def _read_docx(path, filename):
    import docx  # python-docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    # tables too - many reports keep numbers in tables
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    text = '\n'.join(parts).strip()
    return [Document(page_content=text, metadata={'source': filename, 'page': 1})] if text else []


def _read_text(path, filename):
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read().strip()
    return [Document(page_content=text, metadata={'source': filename, 'page': 1})] if text else []


def load_file(path, filename=None):
    """
    Returns a list of Documents (one per PDF page, or one for the whole file).
    Raises ValueError with a friendly message if nothing could be read.
    """
    filename = filename or os.path.basename(path)
    ext = filename.rsplit('.', 1)[-1].lower()

    if ext == 'pdf':
        docs = _read_pdf(path, filename)
    elif ext == 'docx':
        docs = _read_docx(path, filename)
    elif ext in ('txt', 'md', 'csv'):
        docs = _read_text(path, filename)
    else:
        raise ValueError('File type .{} is not supported.'.format(ext))

    if not docs:
        raise ValueError(
            'No text found in {}. If it is a scanned PDF (photos of pages), '
            'it has no selectable text to read.'.format(filename)
        )
    return docs
